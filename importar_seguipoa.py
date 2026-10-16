import os
import hashlib
import json
import re
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
import pymysql
from openpyxl import load_workbook

DB_HOST = "127.0.0.1"
DB_PORT = 3306
DB_USER = "root"
DB_PASSWORD = ""
DB_NAME = "seguipoa_db"

EXCEL_PATH = r"Datos\POA2026SEPTIEMBRE.xlsx"
HOJA = "SEPTIEMBRE 2026"
FILA_ENCABEZADO = 5
FILA_INICIO_DATOS = 6

PERIODO_CORTE = date(2026, 9, 30)


def conexion():
    return pymysql.connect(
        host=DB_HOST,
        port=DB_PORT,
        user=DB_USER,
        password=DB_PASSWORD,
        database=DB_NAME,
        charset="utf8mb4",
        cursorclass=pymysql.cursors.DictCursor,
        autocommit=False,
    )


def limpiar(valor):
    if valor is None:
        return None

    if isinstance(valor, (datetime, date)):
        return valor

    texto = str(valor).strip()
    return texto if texto else None


def a_decimal(valor):
    texto = limpiar(valor)

    if texto is None:
        return None

    if isinstance(valor, (int, float, Decimal)):
        try:
            return Decimal(str(valor))
        except InvalidOperation:
            return None

    texto = str(texto)
    texto = texto.replace("Bs", "").replace(" ", "")
    texto = texto.replace(",", "")

    try:
        return Decimal(texto)
    except (InvalidOperation, ValueError):
        return None


def a_entero(valor):
    numero = a_decimal(valor)

    if numero is None:
        return None

    try:
        return int(numero)
    except (ValueError, TypeError):
        return None


def a_fecha(valor):
    if valor is None:
        return None, None

    if isinstance(valor, datetime):
        return valor.date(), None

    if isinstance(valor, date):
        return valor, None

    texto = limpiar(valor)

    if texto is None:
        return None, None

    # Si contiene varias fechas, se conserva como texto
    if "\n" in texto:
        return None, texto

    texto_limpio = re.sub(r"\s+", " ", texto)

    formatos = [
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%d-%b-%y",
        "%d-%b-%Y",
        "%Y-%m-%d",
    ]

    for formato in formatos:
        try:
            return datetime.strptime(texto_limpio, formato).date(), None
        except ValueError:
            continue

    return None, texto


def hash_archivo(ruta):
    sha256 = hashlib.sha256()

    with open(ruta, "rb") as archivo:
        for bloque in iter(lambda: archivo.read(8192), b""):
            sha256.update(bloque)

    return sha256.hexdigest()


def crear_importacion(conn, ruta):
    with conn.cursor() as cursor:
        cursor.execute(
            """
            INSERT INTO importaciones
                (nombre_archivo, hash_archivo, hoja_origen, gestion,
                 periodo, estado)
            VALUES
                (%s, %s, %s, %s, %s, 'PROCESANDO')
            """,
            (
                os.path.basename(ruta),
                hash_archivo(ruta),
                HOJA,
                2026,
                PERIODO_CORTE,
            ),
        )

        return cursor.lastrowid


def registrar_incidencia(
    cursor,
    importacion_id,
    fila_origen_id,
    columna,
    valor_original,
    codigo,
    descripcion,
    severidad="ADVERTENCIA",
):
    cursor.execute(
        """
        INSERT INTO incidencias_importacion
            (importacion_id, fila_origen_id, columna_origen,
             valor_original, codigo_regla, severidad, descripcion)
        VALUES
            (%s, %s, %s, %s, %s, %s, %s)
        """,
        (
            importacion_id,
            fila_origen_id,
            columna,
            limpiar(valor_original),
            codigo,
            severidad,
            descripcion[:1000],
        ),
    )


def obtener_o_crear_catalogo(cursor, tabla, valor):
    valor = limpiar(valor)

    if valor is None:
        return None

    cursor.execute(
        f"SELECT id FROM {tabla} WHERE nombre = %s",
        (valor,),
    )
    resultado = cursor.fetchone()

    if resultado:
        return resultado["id"]

    cursor.execute(
        f"INSERT INTO {tabla} (nombre) VALUES (%s)",
        (valor,),
    )

    return cursor.lastrowid


def obtener_distrito(cursor, valor):
    texto = limpiar(valor)

    if texto is None:
        return None, None

    cursor.execute(
        "SELECT id FROM distritos WHERE nombre = %s",
        (texto.upper(),),
    )
    resultado = cursor.fetchone()

    if resultado:
        return resultado["id"], texto

    cursor.execute(
        """
        INSERT INTO distritos (codigo, nombre, alcance)
        VALUES (%s, %s, 'SIN_CLASIFICAR')
        """,
        (None, texto.upper()),
    )

    return cursor.lastrowid, texto


def obtener_unidad_ejecutora(cursor, gestion, da, ue, nombre):
    nombre = limpiar(nombre)

    if nombre is None:
        return None, None

    cursor.execute(
        """
        SELECT id FROM unidades_ejecutoras
        WHERE gestion = %s
          AND da_codigo = %s
          AND ue_codigo = %s
          AND nombre = %s
        """,
        (gestion, limpiar(da), limpiar(ue), nombre),
    )
    resultado = cursor.fetchone()

    if resultado:
        return resultado["id"], nombre

    cursor.execute(
        """
        INSERT INTO unidades_ejecutoras
            (gestion, da_codigo, ue_codigo, nombre)
        VALUES (%s, %s, %s, %s)
        """,
        (gestion, limpiar(da), limpiar(ue), nombre),
    )

    return cursor.lastrowid, nombre


def insertar_fila_original(cursor, importacion_id, numero_fila, valores):
    datos = []

    for valor in valores:
        if valor is None:
            datos.append(None)
        elif isinstance(valor, (datetime, date)):
            datos.append(valor.isoformat())
        else:
            datos.append(str(valor))

    cursor.execute(
        """
        INSERT INTO seguipoa_filas_originales
            (importacion_id, numero_fila_excel, datos_originales)
        VALUES (%s, %s, %s)
        """,
        (
            importacion_id,
            numero_fila,
            json.dumps(datos, ensure_ascii=False, default=str),
        ),
    )

    return cursor.lastrowid


def insertar_proyecto(cursor, importacion_id, fila_origen_id, fila):
    gestion = a_entero(fila[0]) or 2026
    sisin = limpiar(fila[3])

    if sisin in ("0", "0.0"):
        registrar_incidencia(
            cursor,
            importacion_id,
            fila_origen_id,
            "SISIN",
            sisin,
            "SISIN_CERO",
            "SISIN registrado como 0; se guarda como NULL.",
        )
        sisin = None

    distrito_id, distrito_texto = obtener_distrito(cursor, fila[7])
    unidad_id, unidad_texto = obtener_unidad_ejecutora(
        cursor, gestion, fila[29], fila[30], fila[31]
    )

    condicion_id = obtener_o_crear_catalogo(
        cursor, "condiciones_poa", fila[32]
    )
    tipo_gasto_id = obtener_o_crear_catalogo(
        cursor, "tipos_gasto", fila[33]
    )

    nombre_proyecto = limpiar(fila[5])

    if nombre_proyecto is None:
        registrar_incidencia(
            cursor,
            importacion_id,
            fila_origen_id,
            "PROYECTO",
            None,
            "NOMBRE_PROYECTO_VACIO",
            "La fila no tiene nombre de proyecto.",
            severidad="ERROR",
        )
        return None

    cursor.execute(
        """
        INSERT INTO proyectos
            (gestion, numero_bd, programa, sisin, actividad,
             nombre_proyecto, localizacion, distrito_id,
             distrito_texto_original, da, ue,
             unidad_ejecutora_id, unidad_ejecutora_texto_original,
             condicion_id, tipo_gasto_id,
             importacion_id, fila_origen_id)
        VALUES
            (%s, %s, %s, %s, %s,
             %s, %s, %s,
             %s, %s, %s,
             %s, %s,
             %s, %s,
             %s, %s)
        """,
        (
            gestion,
            limpiar(fila[1]),
            limpiar(fila[2]),
            sisin,
            limpiar(fila[4]),
            nombre_proyecto,
            limpiar(fila[6]),
            distrito_id,
            distrito_texto,
            limpiar(fila[29]),
            limpiar(fila[30]),
            unidad_id,
            unidad_texto,
            condicion_id,
            tipo_gasto_id,
            importacion_id,
            fila_origen_id,
        ),
    )

    return cursor.lastrowid


def insertar_presupuesto_base(cursor, proyecto_id, importacion_id, fila):
    datos = [
        ("PRESUPUESTO_INICIAL", fila[8]),
        ("PRESUPUESTO_MARZO", fila[9]),
        ("COSTO_TOTAL_PROYECTO", fila[69]),
        ("MONTO_CERTIFICACION_POA", fila[67]),
    ]

    for tipo, valor in datos:
        monto = a_decimal(valor)

        if monto is None:
            continue

        cursor.execute(
            """
            INSERT INTO presupuesto_base
                (proyecto_id, tipo, monto, fecha_referencia, importacion_id)
            VALUES (%s, %s, %s, %s, %s)
            """,
            (proyecto_id, tipo, monto, PERIODO_CORTE, importacion_id),
        )


def insertar_financiamiento(cursor, proyecto_id, importacion_id, fila):
    fuentes = [
        ("COPARTICIPACION_TRIBUTARIA", fila[23]),
        ("RECURSOS_PROPIOS", fila[24]),
        ("REGALIAS_MINERAS", fila[25]),
        ("OTROS_RECURSOS_ESPECIFICOS", fila[26]),
        ("RECURSOS_IDH", fila[27]),
        ("OTROS_RECURSOS", fila[28]),
    ]

    for nombre_fuente, valor in fuentes:
        monto = a_decimal(valor)

        if monto is None:
            continue

        cursor.execute(
            "SELECT id FROM fuentes_financiamiento WHERE nombre = %s",
            (nombre_fuente,),
        )
        fuente = cursor.fetchone()

        if not fuente:
            continue

        cursor.execute(
            """
            INSERT INTO financiamiento_proyecto
                (proyecto_id, fuente_id, monto, periodo, importacion_id)
            VALUES (%s, %s, %s, %s, %s)
            """,
            (
                proyecto_id,
                fuente["id"],
                monto,
                PERIODO_CORTE,
                importacion_id,
            ),
        )


def insertar_ejecucion_mensual(cursor, proyecto_id, importacion_id, fila):
    bloques = [
        ("2026-02-01", 74, 78),
        ("2026-03-01", 79, 83),
        ("2026-04-01", 84, 88),
        ("2026-05-01", 89, 93),
        ("2026-06-01", 94, 98),
        ("2026-07-01", 100, 104),
        ("2026-08-01", 106, 110),
        ("2026-09-01", 112, 116),
    ]

    for periodo, inicio, fin in bloques:
        valores = [a_decimal(fila[indice]) for indice in range(inicio, fin + 1)]

        if all(valor is None for valor in valores):
            continue

        cursor.execute(
            """
            INSERT INTO ejecucion_mensual
                (proyecto_id, periodo, presupuesto_vigente, preventivo,
                 compromiso, devengado, pagado, importacion_id)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
            ON DUPLICATE KEY UPDATE
                presupuesto_vigente = VALUES(presupuesto_vigente),
                preventivo = VALUES(preventivo),
                compromiso = VALUES(compromiso),
                devengado = VALUES(devengado),
                pagado = VALUES(pagado),
                importacion_id = VALUES(importacion_id)
            """,
            (
                proyecto_id,
                periodo,
                valores[0],
                valores[1],
                valores[2],
                valores[3],
                valores[4],
                importacion_id,
            ),
        )

        if valores[4] is not None and valores[3] is not None:
            if valores[4] > valores[3]:
                registrar_incidencia(
                    cursor,
                    importacion_id,
                    None,
                    f"EJECUCION_{periodo}",
                    f"Pagado: {valores[4]}; Devengado: {valores[3]}",
                    "PAGADO_MAYOR_DEVENGADO",
                    "El monto pagado supera al devengado.",
                )


def insertar_seguimiento_fisico(cursor, proyecto_id, importacion_id, fila):
    avance_original = fila[54]
    avance = a_decimal(avance_original)
    avance_texto = None

    if avance is None and limpiar(avance_original) is not None:
        avance_texto = limpiar(avance_original)

        registrar_incidencia(
            cursor,
            importacion_id,
            None,
            "% AVANCE FISICO",
            avance_original,
            "AVANCE_FISICO_NO_NUMERICO",
            "El avance físico contiene texto; se conserva como texto.",
        )

    estado_id = obtener_o_crear_catalogo(
        cursor, "estados_proyecto", fila[52]
    )
    estado_detalle_id = obtener_o_crear_catalogo(
        cursor, "estados_detalle_proyecto", fila[53]
    )

    fecha_solucion, _ = a_fecha(fila[60])

    cursor.execute(
        """
        INSERT INTO seguimiento_fisico
            (proyecto_id, periodo,
             estado_proyecto_id, estado_proyecto_texto_original,
             estado_detalle_id, estado_detalle_texto_original,
             porcentaje_avance_fisico, porcentaje_avance_fisico_texto,
             meta, logrado, observacion,
             descripcion_problema, acciones_a_tomar,
             fecha_solucion_estimada,
             estado_situacion_fisico_obra,
             porcentaje_fisico_acumulado,
             total_numero_etapas, etapa_actual,
             importacion_id)
        VALUES
            (%s, %s,
             %s, %s,
             %s, %s,
             %s, %s,
             %s, %s, %s,
             %s, %s,
             %s,
             %s,
             %s,
             %s, %s,
             %s)
        ON DUPLICATE KEY UPDATE
            estado_proyecto_id = VALUES(estado_proyecto_id),
            estado_proyecto_texto_original = VALUES(estado_proyecto_texto_original),
            estado_detalle_id = VALUES(estado_detalle_id),
            estado_detalle_texto_original = VALUES(estado_detalle_texto_original),
            porcentaje_avance_fisico = VALUES(porcentaje_avance_fisico),
            porcentaje_avance_fisico_texto = VALUES(porcentaje_avance_fisico_texto),
            meta = VALUES(meta),
            logrado = VALUES(logrado),
            observacion = VALUES(observacion),
            descripcion_problema = VALUES(descripcion_problema),
            acciones_a_tomar = VALUES(acciones_a_tomar),
            fecha_solucion_estimada = VALUES(fecha_solucion_estimada),
            estado_situacion_fisico_obra = VALUES(estado_situacion_fisico_obra),
            porcentaje_fisico_acumulado = VALUES(porcentaje_fisico_acumulado),
            total_numero_etapas = VALUES(total_numero_etapas),
            etapa_actual = VALUES(etapa_actual),
            importacion_id = VALUES(importacion_id)
        """,
        (
            proyecto_id,
            PERIODO_CORTE,
            estado_id,
            limpiar(fila[52]),
            estado_detalle_id,
            limpiar(fila[53]),
            avance,
            avance_texto,
            a_decimal(fila[55]),
            a_decimal(fila[56]),
            limpiar(fila[57]),
            limpiar(fila[58]),
            limpiar(fila[59]),
            fecha_solucion,
            limpiar(fila[72]),
            a_decimal(fila[73]),
            a_entero(fila[70]),
            a_entero(fila[71]),
            importacion_id,
        ),
    )


def insertar_contrato(cursor, proyecto_id, importacion_id, fila_origen_id, fila):
    cuce = limpiar(fila[34])
    numero_contrato = limpiar(fila[39])
    empresa = limpiar(fila[40])
    monto_original = fila[42]
    monto = a_decimal(monto_original)

    # Validación de monto de contrato
    LIMITE_MONTO = Decimal("999999999999999.99")

    if monto is not None and monto > LIMITE_MONTO:
        registrar_incidencia(
            cursor,
            importacion_id,
            fila_origen_id,
            "MONTO CONTRATO INICIAL",
            monto_original,
            "MONTO_CONTRATO_INVALIDO",
            "El monto del contrato supera el límite permitido o contiene varios valores concatenados. Se guarda como NULL.",
            severidad="ERROR",
        )
        monto = None

    if all(valor is None for valor in [cuce, numero_contrato, empresa, monto]):
        return None

    cursor.execute(
        """
        INSERT INTO contratos
            (proyecto_id, cuce, modalidad_contratacion,
             numero_contrato, nombre_empresa,
             monto_contrato_inicial, plazo_ejecucion_inicial,
             importacion_id, fila_origen_id)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)
        """,
        (
            proyecto_id,
            cuce,
            limpiar(fila[35]),
            numero_contrato,
            empresa,
            monto,
            limpiar(fila[43]),
            importacion_id,
            fila_origen_id,
        ),
    )

    contrato_id = cursor.lastrowid

    eventos = [
        ("PUBLICACION_SICOES", fila[36]),
        ("APERTURA_SOBRES", fila[37]),
        ("ADJUDICACION", fila[38]),
        ("SUSCRIPCION_CONTRATO", fila[41]),
        ("INICIO_OBRA", fila[44]),
        ("CONCLUSION_OBRA", fila[45]),
        ("RECEPCION_PROVISIONAL", fila[46]),
        ("RECEPCION_DEFINITIVA", fila[47]),
        ("REGISTRO_FORMULARIO_500", fila[48]),
    ]

    for tipo_evento, valor in eventos:
        fecha, fecha_texto = a_fecha(valor)

        if fecha is None and fecha_texto is None:
            continue

        if fecha is None and fecha_texto is not None:
            registrar_incidencia(
                cursor,
                importacion_id,
                fila_origen_id,
                tipo_evento,
                fecha_texto,
                "FECHA_MULTIPLE",
                "La celda contiene varias fechas; se conserva como texto.",
            )

        cursor.execute(
            """
            INSERT INTO contrato_eventos
                (contrato_id, proyecto_id, tipo_evento,
                 fecha, fecha_texto_original,
                 fila_origen_id, importacion_id)
            VALUES (%s, %s, %s, %s, %s, %s, %s)
            """,
            (
                contrato_id,
                proyecto_id,
                tipo_evento,
                fecha,
                fecha_texto,
                fila_origen_id,
                importacion_id,
            ),
        )

    return contrato_id


def insertar_planillas(cursor, proyecto_id, contrato_id, importacion_id, fila_origen_id, fila):
    bloques = [
        ("2026-02-01", 118, 119, 120),
        ("2026-03-01", 121, 122, 123),
        ("2026-04-01", 124, 125, 126),
        ("2026-05-01", 127, 128, 129),
        ("2026-06-01", 130, 131, 132),
        ("2026-07-01", 133, 134, 135),
        ("2026-08-01", 136, 137, 138),
        ("2026-09-01", 139, 140, 141),
    ]

    for periodo, indice_numero, indice_monto, indice_contrato in bloques:
        numero = limpiar(fila[indice_numero])
        monto = a_decimal(fila[indice_monto])
        contrato_texto = limpiar(fila[indice_contrato])

        if numero is None and monto is None and contrato_texto is None:
            continue

        tipo = "SIN_CLASIFICAR"

        if numero:
            numero_upper = numero.upper()

            if "CIERRE" in numero_upper:
                tipo = "CIERRE"
            elif "RETENCION" in numero_upper:
                tipo = "RETENCION"
            elif "AVANCE" in numero_upper:
                tipo = "AVANCE"

        cursor.execute(
            """
            INSERT INTO planillas_mensuales
                (proyecto_id, contrato_id, periodo, tipo_planilla,
                 numero_planilla, monto_devengado, monto_contrato,
                 texto_original, fila_origen_id, importacion_id)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            """,
            (
                proyecto_id,
                contrato_id,
                periodo,
                tipo,
                numero,
                monto,
                None,
                f"{numero or ''} | {contrato_texto or ''}",
                fila_origen_id,
                importacion_id,
            ),
        )

        planilla_id = cursor.lastrowid

        if "\n" in str(numero or "") or "\n" in str(monto or ""):
            registrar_incidencia(
                cursor,
                importacion_id,
                fila_origen_id,
                f"PLANILLA_{periodo}",
                f"{numero} | {monto}",
                "PLANILLA_MULTIPLE",
                "La celda parece contener varias planillas; se guarda el texto original.",
            )

            cursor.execute(
                """
                INSERT INTO planilla_detalles
                    (planilla_id, numero_planilla, monto,
                     contrato_referencia, texto_original,
                     fila_origen_id, importacion_id)
                VALUES (%s, %s, %s, %s, %s, %s, %s)
                """,
                (
                    planilla_id,
                    numero,
                    monto,
                    contrato_texto,
                    f"{numero} | {monto} | {contrato_texto}",
                    fila_origen_id,
                    importacion_id,
                ),
            )


def procesar_fila(cursor, importacion_id, numero_fila, valores):
    fila_origen_id = insertar_fila_original(
        cursor, importacion_id, numero_fila, valores
    )

    proyecto_id = insertar_proyecto(
        cursor, importacion_id, fila_origen_id, valores
    )

    if proyecto_id is None:
        return False

    insertar_presupuesto_base(
        cursor, proyecto_id, importacion_id, valores
    )
    insertar_financiamiento(
        cursor, proyecto_id, importacion_id, valores
    )
    insertar_ejecucion_mensual(
        cursor, proyecto_id, importacion_id, valores
    )
    insertar_seguimiento_fisico(
        cursor, proyecto_id, importacion_id, valores
    )

    contrato_id = insertar_contrato(
        cursor, proyecto_id, importacion_id, fila_origen_id, valores
    )

    insertar_planillas(
        cursor,
        proyecto_id,
        contrato_id,
        importacion_id,
        fila_origen_id,
        valores,
    )

    return True


def importar():
    if not os.path.exists(EXCEL_PATH):
        raise FileNotFoundError(f"No se encontró el archivo: {EXCEL_PATH}")

    print(f"Leyendo: {EXCEL_PATH}")

    libro = load_workbook(
        EXCEL_PATH,
        data_only=True,
        read_only=True,
    )

    if HOJA not in libro.sheetnames:
        raise ValueError(f"No existe la hoja: {HOJA}")

    hoja = libro[HOJA]
    filas = list(hoja.iter_rows(values_only=True))

    datos = filas[FILA_INICIO_DATOS - 1:]
    datos = [fila for fila in datos if fila[0] is not None]

    print(f"Filas detectadas: {len(datos)}")

    conn = conexion()
    importacion_id = crear_importacion(conn, EXCEL_PATH)

    insertadas = 0
    errores = 0

    try:
        for indice, valores in enumerate(datos):
            numero_fila_excel = FILA_INICIO_DATOS + indice

            try:
                with conn.cursor() as cursor:
                    procesada = procesar_fila(
                        cursor,
                        importacion_id,
                        numero_fila_excel,
                        valores,
                    )

                conn.commit()

                if procesada:
                    insertadas += 1

            except Exception as error:
                conn.rollback()
                errores += 1

                print(f"Error en fila {numero_fila_excel}: {error}")

                with conn.cursor() as cursor:
                    registrar_incidencia(
                        cursor,
                        importacion_id,
                        None,
                        "FILA_COMPLETA",
                        str(valores[5]) if len(valores) > 5 else None,
                        "ERROR_IMPORTACION",
                        str(error),
                        severidad="ERROR",
                    )

                conn.commit()

        with conn.cursor() as cursor:
            cursor.execute(
                """
                UPDATE importaciones
                SET
                    filas_leidas = %s,
                    filas_insertadas = %s,
                    filas_con_error = %s,
                    estado = %s
                WHERE id = %s
                """,
                (
                    len(datos),
                    insertadas,
                    errores,
                    "COMPLETADA_CON_ALERTAS" if errores else "COMPLETADA",
                    importacion_id,
                ),
            )

        conn.commit()

        print("Importación finalizada.")
        print(f"ID de importación: {importacion_id}")
        print(f"Filas leídas: {len(datos)}")
        print(f"Filas insertadas: {insertadas}")
        print(f"Filas con error: {errores}")

    finally:
        conn.close()


if __name__ == "__main__":
    importar()